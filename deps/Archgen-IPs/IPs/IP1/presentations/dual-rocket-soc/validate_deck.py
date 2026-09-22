#!/usr/bin/env python3
"""Validate slide evidence, rendered text, and geometry; create review montages."""

from collections import Counter
import hashlib
import json
from pathlib import Path
import re

import fitz
from PIL import Image, ImageDraw
from pptx import Presentation


HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
NAME = "Dual_Rocket_SoC_Overview"


def tokens(value):
    return Counter(re.findall(r"[a-z0-9]+", value.lower()))


def main():
    manifest = json.loads((HERE / "slide-manifest.json").read_text())
    for source, expected in manifest["source_sha256"].items():
        assert hashlib.sha256((ROOT / source).read_bytes()).hexdigest() == expected, source
    prs = Presentation(HERE / f"{NAME}.pptx")
    pdf = fitz.open(HERE / f"{NAME}.pdf")
    assert len(prs.slides) == len(pdf) == len(manifest["slides"]) == 24
    out = HERE / "qa"
    out.mkdir(exist_ok=True)
    result = {"slides": len(pdf), "source_hashes_verified": len(manifest["source_sha256"]), "missing_rendered_tokens": [], "text_overlaps": [], "out_of_page": []}
    thumbs = []
    for i, (slide, page) in enumerate(zip(prs.slides, pdf), 1):
        assert slide.has_notes_slide
        assert "Evidence references" in slide.notes_slide.notes_text_frame.text
        source_text = "\n".join(shape.text for shape in slide.shapes if shape.has_text_frame)
        rendered = page.get_text()
        missing = tokens(source_text) - tokens(rendered)
        if missing:
            result["missing_rendered_tokens"].append({"slide": i, "tokens": dict(missing)})
        lines = []
        for block in page.get_text("dict")["blocks"]:
            for row in block.get("lines", []):
                value = "".join(span["text"] for span in row["spans"])
                if value.strip():
                    box = fitz.Rect(row["bbox"])
                    lines.append((box, value))
                    if not page.rect.contains(box):
                        result["out_of_page"].append({"slide": i, "text": value})
        for j, (a, a_text) in enumerate(lines):
            for b, b_text in lines[j + 1:]:
                overlap = a & b
                if overlap.width > 3 and overlap.height > 3:
                    result["text_overlaps"].append({"slide": i, "text_a": a_text, "text_b": b_text})
        pix = page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False)
        image = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
        image.save(out / f"slide-{i:02d}.png")
        image.thumbnail((960, 540))
        thumbs.append(image)
    for start in range(0, len(thumbs), 8):
        sheet = Image.new("RGB", (1968, 4 * 584 + 24), "#dbe2df")
        draw = ImageDraw.Draw(sheet)
        for offset, im in enumerate(thumbs[start:start + 8]):
            x, y = 16 + (offset % 2) * 984, 20 + (offset // 2) * 584
            sheet.paste(im, (x, y))
            draw.text((x + 4, y + 548), f"SLIDE {start + offset + 1:02d}", fill="#20292d")
        sheet.save(out / f"overview-{start // 8 + 1}.png")
    result["pptx_sha256"] = hashlib.sha256((HERE / f"{NAME}.pptx").read_bytes()).hexdigest()
    result["pdf_sha256"] = hashlib.sha256((HERE / f"{NAME}.pdf").read_bytes()).hexdigest()
    (out / "validation.json").write_text(json.dumps(result, indent=2) + "\n")
    print(json.dumps(result, indent=2))
    assert not result["missing_rendered_tokens"], "Rendered text is missing"
    assert not result["text_overlaps"], "Rendered text overlaps"
    assert not result["out_of_page"], "Rendered text leaves the page"


if __name__ == "__main__":
    main()
