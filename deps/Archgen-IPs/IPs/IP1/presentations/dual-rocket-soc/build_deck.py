#!/usr/bin/env python3
"""Build an editable, evidence-backed overview of the recorded dual-Rocket SoC."""

from collections import Counter
import csv
import hashlib
import json
import os
from pathlib import Path

from PIL import ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
DOC = ROOT / "docs/dual-rocket"
OUTPUT = HERE / "Dual_Rocket_SoC_Overview.pptx"
W, H = 13.333333, 7.5
INK, MUTED, PAPER, WHITE = "20292D", "566469", "F5F8F7", "FFFFFF"
GREEN, MINT, CORAL, GOLD, BLUE = "087F72", "E2F2ED", "CE594B", "AC801D", "3279A3"
LINE, LIGHT, REDBG = "D5DEDB", "EBF0EF", "F9ECE8"
FONT, MONO = "DejaVu Sans", "DejaVu Sans Mono"
CONFIG = "docs/dual-rocket/CONFIGURATION.md"
SOC = "docs/dual-rocket/SOC.md"
SIM = "docs/dual-rocket/reports/simulation/README.md"
LINT = "docs/dual-rocket/reports/lint/README.md"
SYNTH = "docs/dual-rocket/reports/synthesis/README.md"
FLOW = "scripts/dual-rocket/README.md"
WORKFLOW = "docs/dual-rocket/reports/workflow/README.md"
ARTIFACT = "docs/dual-rocket/artifacts/"
CY = "e602d917dcc495c58cabe906535e411707096c9c"
RC = "55bcad0f59436de98ea510334121de8546b9e9d7"
ORFS = "68cc9bc974502b4786a68e9f51a092e0fcb56e82"
UPSTREAM_CONFIG = f"https://github.com/ucb-bar/chipyard/blob/{CY}/generators/chipyard/src/main/scala/config/RocketConfigs.scala#L11-L17"
UPSTREAM_CORE = f"https://github.com/chipsalliance/rocket-chip/blob/{RC}/src/main/scala/rocket/Configs.scala#L15-L60"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(W), Inches(H)
prs.core_properties.title = "Dual-Rocket SoC: Configuration, Architecture and Verification"
prs.core_properties.subject = "Chipyard DualRocketConfig; recorded September 6, 2026 evidence"
prs.core_properties.author = "sainadh812 project | Prepared with Codex"
prs.core_properties.keywords = "Chipyard, Rocket, RISC-V, SoC, lint, synthesis, simulation"
SLIDES = []


def rgb(value):
    return RGBColor.from_string(value)


def rect(s, x, y, w, h, fill=WHITE, stroke=None, width=1):
    shape = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    if stroke:
        shape.line.color.rgb = rgb(stroke)
        shape.line.width = Pt(width)
    else:
        shape.line.fill.background()
    shape._element.spPr.append(OxmlElement("a:effectLst"))
    return shape


def fitted_size(value, w, h, size, font, bold):
    family = "DejaVuSansMono" if font == MONO else "DejaVuSans"
    filename = family + ("-Bold" if bold else "") + ".ttf"
    font_path = str(Path(os.environ.get("DEJAVU_FONT_DIR", "/usr/share/fonts/truetype/dejavu")) / filename)
    width, height = (w - 0.025) * 72, (h - 0.03) * 72
    paragraphs = str(value).split("\n")
    # Measure before writing OOXML; prevent clipping of long identifiers as well as wrapped text.
    for candidate in range(round(size * 2), 15, -1):
        current = candidate / 2
        face = ImageFont.truetype(font_path, round(current * 4))
        length = lambda item: face.getlength(item) / 4
        count, longest = 0, 0
        for paragraph in paragraphs:
            line_text = ""
            count += 1
            for word in paragraph.split():
                longest = max(longest, length(word))
                trial = f"{line_text} {word}" if line_text else word
                if line_text and length(trial) > width:
                    count += 1
                    line_text = word
                else:
                    line_text = trial
        required = count * current * 1.28 + max(0, len(paragraphs) - 1) * 5
        if longest <= width and required <= height:
            return current
    raise ValueError(f"Text cannot fit {w} x {h}: {value}")


def text(s, value, x, y, w, h, size=20, color=INK, bold=False,
         align=PP_ALIGN.LEFT, font=FONT, fill=None):
    size = fitted_size(value, w, h, size, font, bold)
    shape = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.01)
    tf.margin_top = tf.margin_bottom = Inches(0.015)
    tf.vertical_anchor = MSO_ANCHOR.TOP
    if fill:
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(fill)
    for i, line in enumerate(str(value).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.alignment = align
        p.space_before, p.space_after = Pt(0), Pt(5)
        p.line_spacing = 1.10
        p.font.name, p.font.size = font, Pt(size)
        p.font.bold, p.font.color.rgb = bold, rgb(color)
    return shape


def line(s, x1, y1, x2, y2, color=GREEN, width=2, arrow=False):
    shape = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    shape.line.color.rgb = rgb(color)
    shape.line.width = Pt(width)
    if arrow:
        end = OxmlElement("a:tailEnd")
        end.set("type", "triangle")
        shape.line._get_or_add_ln().append(end)
    return shape


def node(s, title, subtitle, x, y, w, h, accent=GREEN, fill=WHITE, size=20):
    rect(s, x, y, w, h, fill, LINE)
    rect(s, x, y, 0.055, h, accent)
    text(s, title, x + 0.16, y + 0.13, w - 0.3, 0.43, size, accent, True)
    text(s, subtitle, x + 0.16, y + 0.63, w - 0.3, h - 0.71, size - 3, MUTED)


def metric(s, number, label, x, y, w=2.85, color=GREEN, note=""):
    text(s, number, x, y, w, 0.77, 39, color, True)
    text(s, label, x, y + 0.89, w, 0.70, 18, INK, True)
    if note:
        text(s, note, x, y + 1.66, w, 0.8, 14, MUTED)


def bullets(s, items, x, y, w, size=21, gap=0.82, color=INK):
    for i, item in enumerate(items):
        rect(s, x, y + i * gap + 0.13, 0.07, 0.07, GREEN)
        text(s, item, x + 0.24, y + i * gap, w - 0.24, gap - 0.08, size, color)


def band(s, label, y=6.23, color=GREEN, fill=MINT, size=17):
    rect(s, 0.55, y, 12.23, 0.55, fill)
    text(s, label, 0.76, y + 0.10, 11.80, 0.34, size, color, True)


def table(s, headers, rows, widths, x=0.65, y=1.78, row_h=0.53, size=16):
    xs = [x]
    for width in widths:
        xs.append(xs[-1] + width)
    rect(s, x, y, sum(widths), row_h, INK)
    for i, header in enumerate(headers):
        text(s, header, xs[i] + 0.12, y + 0.10, widths[i] - 0.22, row_h - 0.12, size - 1, WHITE, True)
    for r, row in enumerate(rows):
        yy = y + (r + 1) * row_h
        rect(s, x, yy, sum(widths), row_h, WHITE if r % 2 == 0 else LIGHT)
        for i, cell in enumerate(row):
            text(s, cell, xs[i] + 0.12, yy + 0.09, widths[i] - 0.22, row_h - 0.10, size, INK, i == 0)


def slide(title, section, subtitle="", sources=(), notes=""):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(PAPER)
    rect(s, 0, 0, W, 0.08, GREEN)
    text(s, section.upper(), 0.58, 0.30, 11.9, 0.25, 10, GREEN, True)
    text(s, title, 0.58, 0.77, 12.15, 0.62, 31, INK, True)
    if subtitle:
        text(s, subtitle, 0.60, 1.43, 12.0, 0.41, 16, MUTED)
    line(s, 0.58, 7.04, 12.75, 7.04, LINE, 0.7)
    text(s, "DUAL ROCKET  /  CHIPYARD  /  EVIDENCE: 06 SEP 2026", 0.59, 7.14, 7.6, 0.18, 8.5, MUTED)
    text(s, f"{len(SLIDES) + 1:02d}", 12.07, 7.11, 0.64, 0.25, 11, GREEN, True, PP_ALIGN.RIGHT)
    local = [p for p in sources if not p.startswith("http")]
    short_refs = []
    for p in local[:2]:
        short_refs.append(p.replace("docs/dual-rocket/", ""))
    if short_refs:
        text(s, "Evidence: " + " | ".join(short_refs), 0.60, 6.84, 12.0, 0.20, 8.2, MUTED)
    note_body = notes + "\n\nEvidence references (relative to project root unless URL):\n" + "\n".join(sources)
    s.notes_slide.notes_text_frame.text = note_body.strip()
    SLIDES.append({"number": len(SLIDES) + 1, "title": title, "section": section, "sources": list(sources), "speaker_notes": notes})
    return s


def build():
    sim = json.loads((DOC / "reports/simulation/summary.json").read_text())
    lint = json.loads((DOC / "reports/lint/evidence/summary.json").read_text())
    syn = json.loads((DOC / "reports/synthesis/mapped-statistics.json").read_text())["design"]
    with (DOC / "reports/simulation/all-results.tsv").open() as f:
        actual = Counter(row["result"] for row in csv.DictReader(f, delimiter="\t"))
    assert actual == sim["unique_results"] == {"PASS": 391, "FAIL": 25, "WALL_TIMEOUT": 6}
    assert syn["num_cells"] == 383763 and syn["num_cells_by_type"]["fakeram45_512x64"] == 194
    assert lint["verilator"]["warnings"] == 2484 and lint["slang"]["warnings"] == 21
    assert lint["verilator"]["errors"] == lint["slang"]["errors"] == 0
    assert syn["num_memories"] == syn["num_processes"] == 0
    assert sum(n for name, n in syn["num_cells_by_type"].items() if name.startswith("DFF")) == 55694
    assert syn["num_cells_by_type"]["DLL_X1"] == 1
    assert len((DOC / "artifacts/rtl.sha256").read_text().splitlines()) == 476

    s = slide("Dual-Rocket SoC", "Project overview", sources=(CONFIG, SOC, SIM, SYNTH), notes="This presentation explains an existing Chipyard-generated dual-Rocket design and the recorded verification work. It is not a silicon product announcement. Created on 2026-09-08 from the September 6 evidence; no EDA or functional simulation was rerun to prepare these slides. All diagrams are logical illustrations, not physical layouts.")
    text(s, "Configuration, architecture\nand verification", 0.62, 1.85, 7.0, 1.60, 34, INK, True)
    text(s, "Two 64-bit RISC-V cores. One shared-memory SoC.\nBuilt from pinned Chipyard sources and evaluated\nwith open-source tools.", 0.66, 3.63, 6.8, 1.29, 22, MUTED)
    node(s, "HART 0", "Rocket + FPU\nPrivate I/D caches", 8.10, 1.89, 2.06, 1.62, GREEN, WHITE, 18)
    node(s, "HART 1", "Rocket + FPU\nPrivate I/D caches", 10.52, 1.89, 2.06, 1.62, GREEN, WHITE, 18)
    line(s, 9.14, 3.51, 9.14, 3.95)
    line(s, 11.55, 3.51, 11.55, 3.95)
    line(s, 9.14, 3.95, 11.55, 3.95)
    line(s, 10.34, 3.95, 10.34, 4.28, arrow=True)
    node(s, "SHARED SYSTEM", "512 KiB L2 + memory fabric\nBoot, debug, interrupts, UART", 8.10, 4.30, 4.48, 1.53, BLUE, WHITE, 19)
    text(s, "Prepared for sainadh812  |  08 September 2026", 0.65, 5.88, 7.0, 0.38, 16, GREEN, True)
    band(s, "Scope: generated RTL and exploratory mapping. No fabricated chip or physical signoff.", 6.29, size=16)

    s = slide("What does this SoC do?", "Purpose", "A SoC combines processors, memory infrastructure and system services.", (SOC, SIM), "Each Rocket hart executes software instructions. The shared coherent hierarchy permits coordinated access to memory, while UART, interrupts and debug provide system services. The demonstrated use is bare-metal RTL simulation. Sv39 and privilege modes are present, but Linux boot was not tested. Application examples are categories of use, not benchmark performance claims.")
    for x, title, sub, color in [(0.65, "COMPUTE", "Two instruction streams\nInteger and floating-point work", GREEN), (4.87, "SHARE DATA", "Caches and coherent interconnect\nMemory and atomic operations", BLUE), (9.09, "CONTROL", "Boot, interrupts, UART\nExternal memory and debug", CORAL)]:
        node(s, title, sub, x, 2.03, 3.62, 1.72, color, WHITE, 21)
    text(s, "Demonstrated in simulation", 0.66, 4.11, 5.6, 0.4, 22, GREEN, True)
    bullets(s, ["RISC-V ISA programs and scalar benchmarks", "Both-hart output, shared data and atomic updates"], 0.69, 4.72, 6.40, 19, 0.63)
    text(s, "Not yet demonstrated", 7.26, 4.11, 5.34, 0.4, 21, MUTED, True)
    text(s, "Custom firmware and multicore applications;\nOS bring-up needs its own software and tests.", 7.28, 4.73, 5.3, 1.04, 20, MUTED)
    band(s, "This is a development and verification platform, not a complete board-ready computer.")

    s = slide("How the project was built", "Process", "The work changed the selected configuration and tooling, not the generated CPU RTL.", (CONFIG, FLOW, WORKFLOW, SIM), "The selected DualRocketConfig already existed upstream. Generation used a prepared, pinned checkout and cached generator JARs. A simulation campaign completed before the documentation-driven lint and synthesis reruns. Configuration and SoC documentation were completed before those fresh reruns. Report packaging preserves failed attempts, diagnostics and provenance. The final generation-wrapper test used an up-to-date cache, not a new elaboration.")
    stages = [("01", "PIN", "Record source revisions\nand tool versions"), ("02", "SELECT", "Choose the existing\nDualRocketConfig"), ("03", "GENERATE", "Elaborate the full SoC\nand emit RTL + metadata"), ("04", "VERIFY", "Run RTL simulation,\nlint and synthesis"), ("05", "REPORT", "Preserve outcomes,\nlimits and input hashes")]
    for i, (n, title, body) in enumerate(stages):
        x = 0.66 + i * 2.53
        text(s, n, x, 2.10, 1.6, 0.67, 36, GREEN, True)
        text(s, title, x, 2.98, 2.20, 0.4, 18, INK, True)
        text(s, body, x, 3.61, 2.19, 1.65, 18, MUTED)
        if i < 4:
            line(s, x + 1.66, 2.46, x + 2.31, 2.46, GREEN, 2, True)
    text(s, "Configuration and SoC documented before fresh lint/synthesis; simulation predates those reruns.", 0.70, 5.65, 11.9, 0.43, 16, MUTED)
    band(s, "Recorded evidence includes unsuccessful attempts and warnings, not only the passing summaries.", size=16)

    s = slide("The exact two-core configuration", "Configuration", "Chipyard already provides this class; no hand-duplication of a Rocket core was needed.", (CONFIG, UPSTREAM_CONFIG, UPSTREAM_CORE), "The CDE configuration composition searches leftmost fragments first; it is parameter composition, not concatenated RTL. WithNHugeCores(2) creates two tile parameter entries and updates NumTiles. AbstractConfig supplies the surrounding SoC; by itself it has no tiles. Huge names an upstream parameter preset and does not measure performance. Private L1 instances are duplicated and the shared directory width changes for the second coherent client.")
    rect(s, 0.66, 2.00, 12.00, 2.12, INK)
    text(s, "class DualRocketConfig extends Config(\n  new freechips.rocketchip.rocket.WithNHugeCores(2) ++\n  new chipyard.config.AbstractConfig)", 0.98, 2.34, 11.35, 1.44, 21, WHITE, font=MONO)
    node(s, "SINGLE-CORE BASELINE", "RocketConfig uses WithNHugeCores(1).", 0.66, 4.43, 5.79, 1.35, BLUE, WHITE, 19)
    node(s, "SELECTED DESIGN", "DualRocketConfig uses WithNHugeCores(2).", 6.77, 4.43, 5.89, 1.35, GREEN, WHITE, 19)
    band(s, "The second hart shares the system fabric and L2; this is not two independent SoCs.", size=17)

    s = slide("From Scala to generated hardware", "Generation", "The configuration is a recipe. Chipyard elaborates that recipe into RTL.", (CONFIG, ARTIFACT + "rtl.sha256"), "chipyard.Generator selects chipyard.harness.TestHarness, containing ChipTop and DigitalTop. Chisel elaborates the design and Diplomacy negotiates buses, devices and clocks. FIRRTL and annotations are lowered by CIRCT firtool. The recorded options include --repl-seq-mem, --split-verilog and --export-module-hierarchy. Logical memory interfaces are not foundry SRAM selections. The 476-file manifest includes simulation resources and is not an instantiated DUT-module count. The native simulator is a separate build target.")
    stages = [("Scala + CDE", "Select parameters"), ("Chisel", "Elaborate + negotiate"), ("FIRRTL", "Intermediate design"), ("firtool", "Lower + split RTL")]
    for i, (title, sub) in enumerate(stages):
        x = 0.68 + i * 3.14
        node(s, title, sub, x, 2.03, 2.58, 1.36, GREEN if i < 2 else BLUE, WHITE, 20)
        if i < 3:
            line(s, x + 2.6, 2.73, x + 3.02, 2.73, GREEN, 2, True)
    text(s, "Generated outputs", 0.69, 3.96, 5.1, 0.42, 23, INK, True)
    bullets(s, ["Split Verilog / SystemVerilog", "DTS, memory map, L2 metadata", "Memory interfaces and test selections"], 0.71, 4.53, 7.00, 19, 0.54)
    metric(s, "476", "generated RTL source files", 8.45, 3.98, 4.0, note="Snapshot includes simulation resources.\nOnly the selected hierarchy is analyzed.")
    band(s, "RTL generation and simulator compilation are separate make targets.")

    s = slide("Inside the full ChipTop SoC", "Architecture", "Simplified functional grouping; individual bus adapters and control branches are omitted.", (SOC, ARTIFACT + "chipyard.harness.TestHarness.DualRocketConfig.l2.json"), "Core data/instruction requests enter the TileLink system bus. The shared inclusive L2 connects SBUS to MBUS. The control-side path is SBUS to CBUS to PBUS, and serial ingress uses FBUS. The services box groups functions for readability and is not an assertion that all services attach to PBUS. Scratchpad is instantiated despite DTS disabled status. External memory is beyond the ChipTop boundary; the AXI window does not imply an on-chip DRAM controller or PHY. This is a logical diagram, not a floorplan; 64-bit AXI describes data width, not address width.")
    rect(s, 0.69, 1.96, 9.06, 4.78, WHITE, LINE, 1.2)
    text(s, "ChipTop / DigitalTop", 0.91, 2.09, 5.0, 0.34, 15, MUTED, True)
    node(s, "ROCKET TILE 0", "RV64 + FPU\n32 KiB I$ + 32 KiB D$", 1.00, 2.60, 3.43, 1.31, GREEN, MINT, 18)
    node(s, "ROCKET TILE 1", "RV64 + FPU\n32 KiB I$ + 32 KiB D$", 5.20, 2.60, 3.43, 1.31, GREEN, MINT, 18)
    line(s, 2.72, 3.91, 2.72, 4.12)
    line(s, 6.92, 3.91, 6.92, 4.12)
    rect(s, 1.00, 4.14, 7.63, 0.36, GREEN)
    text(s, "TileLink SBUS", 2.30, 4.15, 4.90, 0.3, 15, WHITE, True, PP_ALIGN.CENTER)
    line(s, 2.72, 4.50, 2.72, 4.94, arrow=True)
    line(s, 6.92, 4.50, 6.92, 4.94, arrow=True)
    text(s, "CBUS / PBUS control path", 5.14, 4.58, 3.50, 0.28, 12, MUTED, align=PP_ALIGN.CENTER, fill=WHITE)
    node(s, "SHARED L2", "512 KiB inclusive cache", 1.00, 4.97, 3.43, 1.03, BLUE, LIGHT, 17)
    node(s, "SYSTEM SERVICES", "Boot / debug / IRQ / UART", 5.20, 4.97, 3.43, 1.03, CORAL, LIGHT, 17)
    line(s, 2.72, 6.00, 2.72, 6.21)
    rect(s, 1.00, 6.22, 3.43, 0.33, BLUE)
    text(s, "MBUS", 1.15, 6.23, 3.1, 0.27, 14, WHITE, True, PP_ALIGN.CENTER)
    line(s, 4.43, 6.39, 5.15, 6.39, BLUE, 2, True)
    rect(s, 5.20, 6.19, 3.43, 0.40, LIGHT, LINE)
    text(s, "64 KiB scratchpad", 5.35, 6.23, 3.1, 0.28, 15, INK, True)
    line(s, 4.74, 6.39, 4.74, 6.67, BLUE)
    line(s, 4.74, 6.67, 9.25, 6.67, BLUE)
    line(s, 9.25, 6.67, 9.25, 5.65, BLUE)
    line(s, 9.25, 5.65, 10.0, 5.65, BLUE, 2, True)
    node(s, "AXI4 MEMORY", "256 MiB window\n64-bit data\nOutside ChipTop", 10.01, 4.52, 2.66, 1.96, BLUE, WHITE, 17)
    text(s, "Other ports:\nUART and JTAG\nSerial TileLink\nClock and reset", 10.01, 2.58, 2.66, 1.45, 17, MUTED)

    s = slide("What each Rocket tile contains", "Processors", "Hart 0 and hart 1 use the same core preset and private L1 geometry.", (SOC, UPSTREAM_CORE), "The emitted DTS and pinned Rocket configuration confirm these properties. The generated ISA string includes b, but this must not be read as every bit-manipulation extension; Zba/Zbb/Zbs are explicitly selected, not Zbc. Single-instruction decode and retirement are in-order. FPU support includes F, D and Zfh. No V extension is selected. Cache capacities exclude metadata and padding. DTS CPU frequency fields are zero; no achieved CPU frequency is claimed.")
    table(s, ["RESOURCE", "PER-HART CONFIGURATION"], [
        ("Execution", "RV64, in-order, single-instruction decode / retirement"),
        ("Arithmetic", "I/M/A/C + Zba/Zbb/Zbs; floating point F/D/Zfh"),
        ("Privilege and MMU", "Machine / supervisor / user modes; Sv39 virtual memory"),
        ("Instruction cache", "32 KiB, 8 ways, 64 sets, 64-byte lines"),
        ("Data cache", "32 KiB, 8 ways; blocking cache (nMSHRs = 0)"),
        ("Protection and TLBs", "8 PMP entries; 32-entry instruction and data TLBs"),
    ], [3.2, 8.84], y=2.02, row_h=0.54, size=17)
    band(s, "Not selected: vector V, Zbc carry-less multiply, or Zicboz cache-block zero.", size=16)

    s = slide("The memory hierarchy", "Memory", "Architectural capacities are different from mapped SRAM bits and library cell area.", (SOC, SYNTH), "Each hart has 32 KiB instruction and 32 KiB data L1 caches, giving 128 KiB total across two harts. Shared L2 is 512 KiB, eight-way, 1024 sets, 64-byte lines, one coherence bank and seven MSHRs. The 64 KiB scratchpad exists in RTL but its device-tree status is disabled; software must account for that. The external window is 256 MiB at 0x80000000 and needs an external AXI memory subsystem. There is no DDR controller/PHY or 256 MiB on-chip memory.")
    entries = [("PRIVATE L1", "128 KiB total", "32 KiB I$ + 32 KiB D$ per hart", GREEN), ("SHARED L2", "512 KiB", "Inclusive, 8-way, one coherence bank", BLUE), ("SCRATCHPAD", "64 KiB", "On-chip at 0x08000000; DTS disabled", GOLD), ("MAIN MEMORY", "256 MiB window", "External AXI4 region at 0x80000000", CORAL)]
    for i, (title, capacity, body, color) in enumerate(entries):
        yy = 2.01 + i * 0.94
        rect(s, 0.67, yy, 2.59, 0.78, color)
        text(s, title, 0.85, yy + 0.22, 2.2, 0.32, 16, WHITE, True)
        text(s, capacity, 3.63, yy + 0.12, 3.25, 0.54, 25, color, True)
        text(s, body, 7.10, yy + 0.19, 5.34, 0.47, 18, MUTED)
    band(s, "External main memory is not on-chip DRAM. ChipTop has no DDR controller or DDR PHY.", size=16)

    s = slide("Boot, peripherals and external interfaces", "System services", "The base SoC includes essential control and debug, not every optional Chipyard device.", (SOC, CONFIG), "CLINT provides timer/software interrupts to both harts. PLIC has one device source (UART) and four M/S interrupt contexts across two physical harts. JTAG debug supports system-bus access. Boot ROM has a 64 KiB address window, not necessarily 64 KiB of program bytes. ChipTop exposes external clocks and reset; the clock generator is passthrough, not a PLL. Decoupled serial TileLink uses 32-bit phits/flits; it is not UART. The current SoC has no FFT or other RoCC accelerator.")
    node(s, "BOOT + INTERRUPTS", "Boot ROM / boot selection\nCLINT timer + software IRQ\nPLIC: UART device interrupt", 0.66, 2.02, 3.72, 2.55, GREEN, WHITE, 19)
    node(s, "EXTERNAL INTERFACES", "UART RX/TX and JTAG\nAXI4 external memory\nSerial TileLink; clocks/reset", 4.80, 2.02, 3.72, 2.55, BLUE, WHITE, 19)
    node(s, "NOT IN THIS CONFIG", "FFT / RoCC / vector unit\nGPIO, I2C, SPI flash\nEthernet or DDR controller", 8.94, 2.02, 3.72, 2.55, CORAL, WHITE, 19)
    text(s, "Do we have FFT?", 0.68, 5.00, 4.07, 0.5, 27, CORAL, True)
    text(s, "No dedicated FFT hardware. A Rocket CPU can execute FFT software;\nan FFT accelerator needs a new integration and its own verification.", 4.80, 4.99, 7.86, 1.00, 20, INK)
    band(s, "Optional generator code in Chipyard does not mean that block exists in this generated SoC.", size=16)

    s = slide("Key software-visible address regions", "Memory map", "Addresses come from the emitted memory map and device tree.", (SOC, ARTIFACT + "chipyard.harness.TestHarness.DualRocketConfig.memmap.json"), "This slide is a selected address map, not all 12 regions. The full map also contains boot-address, error-response, tile-gating and tile-reset windows. End addresses in the source report are inclusive. Register windows are address apertures, not RAM capacity. Scratchpad status is disabled in DTS despite being instantiated. PLIC's 64 MiB aperture is not 64 MiB of interrupt storage. Debug is at zero, not ordinary RAM.")
    table(s, ["BASE ADDRESS", "WINDOW", "RESOURCE / PURPOSE"], [
        ("0x00000000", "4 KiB", "Debug module"), ("0x00010000", "64 KiB", "Boot ROM address window"),
        ("0x02000000", "64 KiB", "CLINT: timer / software interrupts"), ("0x02010000", "4 KiB", "Shared L2 control registers"),
        ("0x08000000", "64 KiB", "On-chip scratchpad"), ("0x0c000000", "64 MiB", "PLIC register aperture"),
        ("0x10020000", "4 KiB", "UART registers"), ("0x80000000", "256 MiB", "External main-memory window"),
    ], [3.2, 2.0, 6.84], y=1.98, row_h=0.44, size=16)
    band(s, "Register-window size is not storage capacity. Full address table is in SOC.md.", size=16)

    s = slide("How two cores cooperate", "Multicore behavior", "Separate instruction streams can coordinate through shared memory and atomic operations.", (SOC, SIM), "This diagram is conceptual, not the instruction-by-instruction sequence of the test. The explicit two-hart atomic smoke requires both harts, peer-data visibility and a total of 256 atomic increments. Two-hart hello requires two hart-identifying outputs. Corrected multicore tests aggregate both hart results before success. Most physical ISA tests park the second hart; ordinary benchmarks can use one-hart startup code. Shared coherent caches do not automatically parallelize a program and atomic checks are not exhaustive coherence verification.")
    node(s, "HART 0", "Runs its own software\nPublishes / checks shared data", 0.68, 2.23, 3.50, 1.65, GREEN, WHITE, 21)
    node(s, "SHARED STATE", "Coherent memory\nAtomic counter updates", 4.92, 2.23, 3.50, 1.65, BLUE, WHITE, 21)
    node(s, "HART 1", "Runs its own software\nPublishes / checks peer data", 9.16, 2.23, 3.50, 1.65, GREEN, WHITE, 21)
    line(s, 4.21, 3.08, 4.81, 3.08, GREEN, 2.3, True)
    line(s, 9.12, 3.08, 8.53, 3.08, GREEN, 2.3, True)
    metric(s, "2 harts", "required by explicit smoke", 0.72, 4.43, 3.8)
    metric(s, "256", "total atomic increments", 4.94, 4.43, 3.8, BLUE)
    metric(s, "PASS", "peer-data visibility checked", 9.19, 4.43, 3.48)
    band(s, "Software must use both harts; two cores do not automatically make every program parallel.", size=16)

    s = slide("Hardware boundary vs simulation harness", "Scope", "Full-SoC synthesis means ChipTop, not the whole simulated environment.", (SOC, SIM, LINT, SYNTH), "TestDriver and TestHarness instantiate the complete ChipTop plus test infrastructure. Simulation uses original behavioral memories and DRAMSim2. +loadmem writes the simulator memory model to load a program; it is not a hardware memory-write port. Simulated JTAG, serial host, UART output and clock sources are not ASIC hardware. Lint/synthesis select the ChipTop hierarchy. Lint uses original behavioral memories; synthesis applies validated SRAM adapters. RTL simulation does not validate the mapped netlist.")
    rect(s, 0.69, 2.05, 11.96, 3.63, LIGHT, LINE)
    text(s, "TestDriver / TestHarness: simulation environment", 0.93, 2.23, 11.38, 0.4, 20, MUTED, True)
    node(s, "ChipTop / DigitalTop", "Both Rocket tiles, L1s and L2\nScratchpad, buses and peripherals\nClock/reset logic and I/O cells", 1.04, 2.99, 5.58, 2.28, GREEN, WHITE, 23)
    node(s, "EXTERNAL TEST MODELS", "DRAMSim2 / simulated memory\nSimJTAG, serial host, UART output\nTest clocks and fast ELF loading", 7.02, 2.99, 5.27, 2.28, CORAL, WHITE, 20)
    band(s, "Synthesis boundary = ChipTop. +loadmem is a simulation shortcut, not fabricated hardware.", size=16)

    s = slide("Three checks answer different questions", "Verification strategy", "Common source evidence; separate methods, outputs and acceptance boundaries.", (SIM, LINT, SYNTH, FLOW), "The completed software simulation campaign predates the fresh documentation-driven lint and synthesis runs. Lint and synthesis can run independently against the same preserved generated source snapshot. Their dates and artifacts are retained separately. Simulation asks whether specific tests execute successfully. Lint checks static diagnostics, not full behavior. Synthesis checks mapping and linkability into libraries, not physical signoff. 476 hashes refer to a full generated-source inventory, not 476 DUT instances.")
    for x, title, question, body, color in [
        (0.66, "SIMULATION", "Does this test execute?", "Verilator + DRAMSim2\nRISC-V programs and multicore tests\nReference checks with Spike", GREEN),
        (4.80, "LINT", "What does static analysis flag?", "Verilator + slang\nFull-ChipTop diagnostics\nWarning baseline and source checks", BLUE),
        (8.94, "SYNTHESIS", "Can the design be mapped?", "Yosys / slang + ABC\nSRAM adapters and libraries\nOpenROAD netlist linking", CORAL),
    ]:
        rect(s, x, 2.08, 3.72, 3.52, WHITE, LINE)
        rect(s, x, 2.08, 3.72, 0.08, color)
        text(s, title, x + 0.19, 2.37, 3.34, 0.42, 22, color, True)
        text(s, question, x + 0.19, 3.07, 3.34, 0.78, 20, INK, True)
        text(s, body, x + 0.19, 4.13, 3.34, 1.22, 17, MUTED)
    band(s, "All three are useful. None alone establishes functional completeness or manufacturing readiness.", size=15.5)

    s = slide("Simulation: configured tests passed", "Results | simulation", "Recorded full-SoC RTL campaign, completed 06 September 2026. Not rerun for this deck.", (SIM, "docs/dual-rocket/reports/simulation/summary.json"), "All 335 configured ISA and 12 configured benchmarks passed. The configured ISA suite is a subset of the 349-program ISA inventory. The entire broadened campaign contains 422 unique programs/variants, with 391 reported PASS, 25 FAIL and six host wall timeouts. Retries are not added to the unique total. PASS results for some legacy programs carry oracle caveats explained on the next slide and in the report. Verilator simulation version is 5.022, distinct from the newer lint executable. Fast ELF loading bypassed serial program transfer.")
    metric(s, "335 / 335", "configured ISA tests", 0.70, 2.00, 4.30)
    metric(s, "12 / 12", "configured benchmarks", 5.08, 2.00, 4.30)
    text(s, "Broader inventory: 422 unique programs / variants", 0.70, 4.18, 11.94, 0.5, 24, INK, True)
    total_w, xx = 11.96, 0.70
    for count, color in [(actual["PASS"], GREEN), (actual["FAIL"], CORAL), (actual["WALL_TIMEOUT"], GOLD)]:
        ww = total_w * count / sim["unique_tests"]
        rect(s, xx, 4.98, ww, 0.46, color)
        xx += ww
    for x, label, color in [(0.72, "391 reported PASS", GREEN), (5.15, "25 FAIL", CORAL), (8.73, "6 wall timeouts", GOLD)]:
        text(s, label, x, 5.59, 3.84, 0.44, 22, color, True)
    band(s, "The expanded inventory did not pass completely. Do not report 422/422 passing.", size=17)

    s = slide("Interpreting failures, timeouts and passes", "Results | simulation", "Test intent and the correctness of the test itself matter.", (SIM, "docs/dual-rocket/reports/simulation/reference-comparison/summary.json"), "Nine non-passing additional ISA probes concern unsupported Zbc/Zicboz or misaligned-data expectations. They are not failures of the configured 335-test subset. The 22 legacy non-passes have 32-by-32 matrix assumptions against a 16-by-16 dataset; some reported legacy passes also have bounds or two-hart result-aggregation defects. Three separately named corrected multicore supplements passed RTL and Spike; original results remain preserved. Matched Spike has 391 PASS/PASS and 31 non-PASS/non-PASS outcomes, not identical exit behavior or cycle-level equivalence. The six final timeouts are host wall limits, not six proven core deadlocks. PMP passed after a bounded extended retry, and a C++ hello parser issue was corrected.")
    node(s, "9 ISA NON-PASSES", "Additional unsupported-feature or\nmisaligned-access expectation probes", 0.68, 2.03, 5.79, 1.52, CORAL, WHITE, 20)
    node(s, "22 LEGACY NON-PASSES", "Matrix geometry / test-oracle defects;\n32 x 32 assumptions vs 16 x 16 data", 6.87, 2.03, 5.79, 1.52, GOLD, WHITE, 20)
    bullets(s, ["Some reported legacy passes also have bounds or result-checking caveats.", "Three corrected multicore supplements passed both RTL and Spike.", "Spike agreed on 391 PASS and 31 non-PASS classifications."], 0.74, 4.00, 11.85, 20, 0.61)
    band(s, "Agreement with Spike is supporting evidence, not formal equivalence or exhaustive verification.", size=16)

    s = slide("Lint: zero errors, warnings retained", "Results | lint", "Both tools analyzed full ChipTop using the original behavioral memories.", (LINT, "docs/dual-rocket/reports/lint/evidence/summary.json"), "Fresh wrappers started both lint tools at 2026-09-06T13:19:52Z. Verilator 5.051 reported 2484 warnings; slang 11.0.448 reported 21 arith-in-shift warnings and 21 notes. Both tool and wrapper exits are zero. All warning-category deltas are zero against the earlier DualRocketConfig baseline, not the historical one-core report. Verilator -Wno-fatal retains warnings but allows completion. No generated RTL edits or new suppressions were introduced. Four SYNCASYNCNET warnings require reset/clock review. EICG_wrapper's latch is consistent with an intentional clock-gate latch, not physical signoff. 419 source files are in slang's elaborated dependencies; the shared 476-file fingerprint is the full generated inventory.")
    table(s, ["TOOL", "ERRORS", "WARNINGS", "NOTES"], [("Verilator 5.051", "0", f"{lint['verilator']['warnings']:,}", "0"), ("slang 11.0.448", "0", "21", "21")], [5.14, 2.05, 2.76, 2.09], y=2.03, row_h=0.64, size=20)
    text(s, "Dominant baseline", 0.70, 4.39, 5.60, 0.39, 21, GREEN, True)
    text(s, "2,047 unused-signal warnings\n424 empty pin connections", 0.72, 4.96, 5.65, 0.91, 20, MUTED)
    text(s, "Review items still visible", 6.88, 4.39, 5.70, 0.39, 21, CORAL, True)
    text(s, "4 reset-related SYNCASYNCNET warnings\n1 clock-gate latch; 21 shift expressions", 6.90, 4.96, 5.70, 0.91, 19, MUTED)
    band(s, "Zero errors is not warning-free signoff. Dedicated CDC/RDC review remains necessary.", size=16)

    s = slide("How full-SoC synthesis was performed", "Implementation flow", "Yosys performs synthesis and mapping; OpenROAD checks the mapped netlist can link.", (SYNTH, "scripts/dual-rocket/synthesis/README.md"), "The flow preserves original RTL and uses an independent output directory. It checks full-ChipTop reachability and SRAM metadata; validates replacement adapters against behavioral memories; reads SystemVerilog with the slang Yosys plugin; maps logic with Yosys/ABC into Nangate45 cells and SRAM into fakeram45 macros; rejects unresolved cells/memories/processes; and links the complete mapped netlist in OpenROAD with Liberty/LEF. The first attempt failed because literal quotes reached Slang filename arguments, after SRAM tests. Relative input symlinks fixed the tooling issue and the full run was repeated. OpenROAD linking was not placement, routing or static timing analysis.")
    items = [("01", "Validate RTL + SRAM inventory", "Check dimensions, masks, instances and source hashes."), ("02", "Check SRAM adapters", "Compare generated memory behavior over seven interfaces."), ("03", "Synthesize and technology-map", "Yosys + slang + ABC; Nangate45 / fakeram45 libraries."), ("04", "Link and audit the result", "OpenROAD reads all mapped instances; recheck inputs.")]
    for i, (n, title, body) in enumerate(items):
        yy = 2.00 + 0.98 * i
        text(s, n, 0.72, yy, 0.73, 0.55, 27, GREEN, True)
        text(s, title, 1.71, yy, 10.79, 0.45, 21, INK, True)
        text(s, body, 1.73, yy + 0.49, 10.70, 0.38, 17, MUTED)
    band(s, "Technology mapping completed. No timing constraint, floorplan, placement or routing was run.", size=16)

    s = slide("SRAM adaptation was a separate task", "Memory implementation", "Logical memories must be matched to available macro width, depth and write-mask behavior.", (SYNTH, "docs/dual-rocket/reports/synthesis/inventory.json"), "Seven distinct generated SRAM interfaces represent 16 logical full-SoC bulk-memory instances. The adapter maps them into 194 fakeram45_512x64 macros (512 words by 64 bits each). Logical capacity is 5,958,656 bits; physical macro capacity is 6,356,992 bits, including padding and mapping overhead. This is not just architectural cache data. The dual-core directory is 1024 by 144 with 18-bit mask granularity; the one-core directory was 1024 by 136 with 17-bit granularity. Wrong widths, mask lanes and a missing instance were rejected by metadata negative controls. All seven interface tests passed three seeds and 212,364 compared cycles. These are finite initialized-memory comparisons, not formal equivalence, X-state proof or full mapped-SoC simulation.")
    metric(s, "7", "interface shapes", 0.72, 2.08, 3.64)
    metric(s, "16", "logical memory instances", 4.95, 2.08, 3.64, BLUE)
    metric(s, "194", "512 x 64 SRAM macros", 9.20, 2.08, 3.43, CORAL)
    text(s, "Two-core-specific change", 0.74, 4.28, 5.7, 0.42, 22, INK, True)
    text(s, "L2 directory: 1024 x 144 bits\nMask granularity: 18 bits", 0.76, 4.91, 5.75, 0.90, 21, MUTED)
    text(s, "Adapter comparison result", 6.89, 4.28, 5.7, 0.42, 22, GREEN, True)
    text(s, "7 interfaces x 3 seeds\n212,364 compared cycles: PASS", 6.91, 4.91, 5.65, 0.90, 21, GREEN)
    band(s, "Finite initialized-state simulation passed. Formal and X-state equivalence were not established.", size=15.5)

    s = slide("Synthesis results: mapped and linked", "Results | synthesis", "Exploratory Nangate45 / fakeram45 library estimate. Not die area or manufacturing signoff.", (SYNTH, "docs/dual-rocket/reports/synthesis/mapped-statistics.json"), "The fresh run completed September 6, 2026. Yosys maps 383,569 standard cells including 55,694 flip-flops and one DLL_X1 latch, plus 194 SRAM macros for 383,763 total instances. No unresolved internal cells, Yosys memories or processes remain. OpenROAD independently links every instance. Total Liberty cell area is 4,004,732.942001 square micrometers; SRAM subtotal is 3,356,478.972 and standard-cell subtotal 648,253.970001. Divide by one million for square millimeters. These are sums of library areas without a floorplan, not die/core area, utilization or timing. The Nangate45/fakeram45 platform is exploratory and non-manufacturable.")
    metric(s, f"{syn['num_cells']:,}", "total mapped instances", 0.71, 1.98, 4.18)
    metric(s, "383,569", "standard cells", 5.05, 1.98, 3.93, BLUE)
    metric(s, "194", "SRAM macros", 9.46, 1.98, 3.15, CORAL)
    area = syn["area"] / 1_000_000
    sram_area = 194 * 17301.438 / 1_000_000
    cell_area = area - sram_area
    text(s, f"{area:.6f} mm\u00b2", 0.74, 4.07, 5.1, 0.77, 35, INK, True)
    text(s, "total Liberty cell-area sum", 6.0, 4.28, 6.3, 0.42, 21, MUTED)
    rect(s, 0.76, 5.01, 11.82 * sram_area / area, 0.47, CORAL)
    rect(s, 0.76 + 11.82 * sram_area / area, 5.01, 11.82 * cell_area / area, 0.47, BLUE)
    text(s, f"SRAM: {sram_area:.6f} mm\u00b2 ({100 * sram_area / area:.1f}%)", 0.78, 5.62, 7.2, 0.39, 18, CORAL, True)
    text(s, f"Standard cells: {cell_area:.6f} mm\u00b2", 8.25, 5.62, 4.38, 0.39, 17, BLUE, True)
    band(s, "55,694 flip-flops + 1 latch are included in standard cells. Unmapped cells / memories / processes: 0.", size=14.5)

    s = slide("What is not proven yet", "Limits and risks", "Successful elaboration, tests and mapping are milestones, not chip signoff.", (SOC, SIM, LINT, SYNTH), "The nominal 500 MHz bus configuration and 500 kHz DTS timebase are not measured silicon clocks. CPU clock-frequency fields are zero. No timing constraints, STA closure, placement, routing, power analysis, DRC/LVS or manufacturability claim was established. Linux boot, full JTAG/OpenOCD flows, exhaustive coherence, coverage percentages, formal verification and SRAM-mapped gate-level regression remain untested. Serial TileLink's selected decoupled PHY has a retained heavy-traffic deadlock warning; the generator recommends a credited PHY when deadlock freedom is required. Prior fast memory loading does not test sustained serial safety. Lint reset warnings need dedicated review.")
    table(s, ["AREA", "REMAINING WORK / BOUNDARY"], [
        ("Clock and timing", "500 MHz is nominal intent, not achieved CPU frequency."),
        ("Physical design", "No placement, routing, power analysis or DRC/LVS."),
        ("Functional closure", "No exhaustive coherence, formal proof or coverage percentage."),
        ("Software / debug", "Linux boot and full JTAG/OpenOCD workflows not tested."),
        ("Serial and reset", "Heavy-traffic serial-PHY warning; reset-domain review needed."),
        ("Mapped behavior", "No full-SoC SRAM-mapped gate-level regression."),
    ], [3.18, 8.86], y=2.01, row_h=0.56, size=17)
    band(s, "Report the design as a verified experiment with explicit limits, not a production-ready ASIC.", color=CORAL, fill=REDBG, size=16)

    s = slide("How to try the next SoC configuration", "Next experiments", "Proposed work only. None of these changes is included in the current generated SoC.", (CONFIG, SOC, FLOW), "This is a proposed engineering sequence, not completed integration. For an FFT accelerator, first initialize and inspect the optional FFT generator at the pinned revision and study its integration rather than assuming the class exists in the main tree. Choose a documented attachment, addressing and interrupt scheme as applicable. New core/cache configurations change memory shapes and need refreshed adapters. Preserve the current experiment, use distinct config names and output roots, regenerate metadata, add feature-specific tests, and rerun lint/synthesis. Quantitative comparisons require common tools and constraints. No FFT hardware or Linux boot is claimed in the current design.")
    node(s, "OPTION A: CORE / CACHE", "Change core count or cache geometry.\nRecheck coherence clients and SRAM shapes.", 0.67, 2.05, 5.79, 1.66, GREEN, WHITE, 19)
    node(s, "OPTION B: FFT ACCELERATOR", "Integrate the optional generator.\nDefine interface and numerical tests.", 6.85, 2.05, 5.79, 1.66, BLUE, WHITE, 19)
    bullets(s, ["Use a new named Config and preserve this baseline.", "Regenerate RTL, device tree, memory map and SRAM inventory.", "Add targeted software tests, then rerun lint and full-SoC synthesis."], 0.76, 4.20, 11.75, 20, 0.63)
    band(s, "Choose the workload first; compare new designs with the same tools and explicit acceptance criteria.", size=15.5)

    s = slide("Where the reproducible process lives", "Reproduction", "The results repository is not a complete toolchain or a vendored Chipyard checkout.", (CONFIG, FLOW, "docs/dual-rocket/README.md"), "Run the commands from the results repository after provisioning the pinned Chipyard checkout and the documented EDA tools. CHIPYARD_ROOT, Java/RISC-V/PATH, OSS_CAD_SUITE_ROOT and ORFS_ROOT must be set as documented. generate.sh can reuse cached upstream output and requires exact source-manifest agreement. lint and synthesis use separate output roots and must not overwrite recorded attempts. The simulation evidence exporter only packages a completed campaign; it does not launch a simulator. The historical top-level rtl directory is single-core, while the dual-core hardware is identified by the generated-src DualRocketConfig path and its preserved hashes. Reports do not vendor raw large outputs, external libraries or binaries.")
    rect(s, 0.69, 2.04, 12.0, 1.94, INK)
    text(s, "bash scripts/dual-rocket/generate.sh\nbash scripts/dual-rocket/lint/run.sh\nbash scripts/dual-rocket/synthesis/run.sh", 0.97, 2.30, 11.45, 1.44, 22, WHITE, font=MONO)
    text(s, "Architecture + process", 0.73, 4.35, 5.88, 0.37, 20, GREEN, True)
    text(s, "docs/dual-rocket/CONFIGURATION.md\ndocs/dual-rocket/SOC.md", 0.75, 4.94, 6.13, 0.93, 16, MUTED, font=MONO)
    text(s, "Evidence + portable scripts", 7.08, 4.35, 5.48, 0.37, 20, BLUE, True)
    text(s, "docs/dual-rocket/reports/\nscripts/dual-rocket/README.md", 7.10, 4.94, 5.45, 0.93, 16, MUTED, font=MONO)
    band(s, "Important: the top-level rtl/ folder is historical single-core RTL, not this dual-core build.", size=16)

    s = slide("Pinned tools and source provenance", "Appendix | versions", "Versions describe the recorded experiment; this is not a recommendation to use arbitrary latest builds.", (CONFIG, SIM, LINT, "docs/dual-rocket/reports/synthesis/tool-versions.txt", SYNTH, WORKFLOW), "Exact full source revisions and dependency hashes are in the local references. Chipyard e602d917dcc495c58cabe906535e411707096c9c; Rocket Chip 55bcad0f59436de98ea510334121de8546b9e9d7; ORFS 68cc9bc974502b4786a68e9f51a092e0fcb56e82. Verilator simulator is 5.022, lint is 5.051 devel v5.050-312-gb1c06fdb0 (mod). slang is 11.0.448+e222e7dc0. Yosys reports 0.68+195 with a dirty suffix; native executable/plugin hashes are retained. OpenROAD v2.0-17598-ga008522d8. Icarus is version 14.0 devel s20260301-403-g5ab23063f-dirty. 500 synthesis input hashes were revalidated unchanged. Generation was not a clean dependency bootstrap and the later wrapper check was cached.")
    table(s, ["LAYER", "RECORDED INPUT"], [
        ("Chipyard / Rocket Chip", "e602d917dcc4 / 55bcad0f5943 (full SHA in notes)"),
        ("Generator stack", "Chisel 6.7.0; Scala 2.13.16; firtool 1.75.0"),
        ("Java", "OpenJDK 20.0.2-internal"),
        ("RTL simulation", "Verilator 5.022; DRAMSim2; matched Spike"),
        ("Lint", "Verilator 5.051 devel; slang 11.0.448"),
        ("Mapping / memory checks", "Yosys 0.68+195; ABC; Icarus Verilog 14.0 devel"),
        ("Netlist link / libraries", "OpenROAD v2.0-17598; pinned ORFS Nangate45 / fakeram45"),
    ], [3.5, 8.54], y=2.00, row_h=0.51, size=16)
    band(s, "476 RTL source hashes preserved; all 500 captured synthesis input hashes rechecked unchanged.", size=15.5)

    s = slide("Quick glossary and evidence guide", "Appendix | reading guide", "All slides have speaker notes with project-relative source references.", (SOC, CONFIG, SIM, LINT, SYNTH), "Glossary: SoC is a system on chip. A hart is a hardware execution thread; this design has one hart per Rocket core. RTL describes digital hardware behavior. TileLink is the on-chip interconnect protocol used here; AXI4 is the external-memory interface. L1 is private to a tile; L2 is shared. CLINT supplies timer/software interrupts; PLIC routes device interrupts. Lint, simulation, synthesis and physical signoff are distinct. The companion slide manifest and speaker-notes Markdown contain all references and SHA-256 fingerprints of the cited local source files. Read the detailed reports for exact test identities, warnings, provenance and limitations. This deck is a derivative explanation and does not replace the raw evidence.")
    table(s, ["TERM", "MEANING IN THIS PROJECT"], [
        ("SoC / hart", "Integrated system / one hardware execution thread"),
        ("RTL / Config", "Hardware description / generator parameter composition"),
        ("L1 / L2", "Private per-core caches / shared inclusive cache"),
        ("TileLink / AXI4", "On-chip fabric / external-memory interface"),
        ("CLINT / PLIC", "Timer + software interrupts / device interrupts"),
    ], [3.18, 8.86], y=2.01, row_h=0.55, size=17)
    text(s, "Start reading: docs/dual-rocket/README.md", 0.77, 5.74, 11.67, 0.5, 23, GREEN, True)
    band(s, "Bottom line: two Rocket cores, one complete SoC, reproducible evidence and clearly stated limits.", size=15.5)

    assert len(prs.slides) == 24
    check_bounds()
    prs.save(OUTPUT)
    sources = sorted({p for item in SLIDES for p in item["sources"] if not p.startswith("http")})
    hashes = {}
    for source in sources:
        path = ROOT / source
        assert path.is_file(), source
        hashes[source] = hashlib.sha256(path.read_bytes()).hexdigest()
    manifest = {"presentation": OUTPUT.name, "prepared_date": "2026-09-08", "evidence_date": "2026-09-06", "slides": SLIDES, "source_sha256": hashes}
    (HERE / "slide-manifest.json").write_text(json.dumps(manifest, indent=2) + "\n")
    notes = ["# Dual-Rocket SoC: Speaker Notes", "", "Prepared 2026-09-08 from recorded 2026-09-06 evidence. No new EDA run.", ""]
    for item in SLIDES:
        notes.extend([f"## {item['number']:02d}. {item['title']}", "", item["speaker_notes"], "", "Sources:", ""])
        for source in item["sources"]:
            target = source if source.startswith("http") else "../../" + source
            notes.append(f"- [{source}]({target})")
        notes.append("")
    (HERE / "speaker-notes.md").write_text("\n".join(notes))
    print(f"Created {OUTPUT.name}: {len(prs.slides)} editable slides; {len(hashes)} referenced source hashes.")


def check_bounds():
    for n, s in enumerate(prs.slides, 1):
        for shape in s.shapes:
            for value in (shape.left, shape.top, shape.width, shape.height):
                assert value >= 0, (n, shape.name, "negative geometry")
            assert shape.left + shape.width <= prs.slide_width + 100, (n, shape.name, "right overflow")
            assert shape.top + shape.height <= prs.slide_height + 100, (n, shape.name, "bottom overflow")


if __name__ == "__main__":
    build()
